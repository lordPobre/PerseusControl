"""Perseus v4 — views/reportes.py
Exportación Excel (RRHH, DT, Remuneraciones) y PDF trabajador.
"""
import hashlib
import logging
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from django.contrib.auth.decorators import login_required
from django.db.models import Min, Max
from django.http import HttpResponse
from django.shortcuts import render, redirect
from django.template.loader import render_to_string
from django.utils import timezone
from datetime import timedelta, time

from apps.asistencia.models import (
    Marcacion, Feriado, Vacacion, LicenciaMedica, Perfil
)
from apps.asistencia.views.mixins import requiere_empleador, requiere_empresa

logger = logging.getLogger(__name__)


@login_required
@requiere_empleador
@requiere_empresa
def exportar_excel_empresa(request):
    """Excel detallado de marcaciones para RRHH."""
    empresa = request.empresa
    marcas  = (
        Marcacion.objects
        .filter(empresa=empresa)
        .select_related('trabajador', 'trabajador__perfil')
        .order_by('trabajador', 'timestamp')
    )

    fi = request.GET.get('fecha_inicio')
    ff = request.GET.get('fecha_fin')
    if fi and ff:
        marcas = marcas.filter(timestamp__date__range=[fi, ff])

    # Agrupar por (trabajador, día)
    reporte = {}
    for m in marcas:
        fl  = timezone.localtime(m.timestamp)
        key = (m.trabajador_id, fl.date())
        if key not in reporte:
            p = getattr(m.trabajador, 'perfil', None)
            reporte[key] = {
                'fecha': fl.date(),
                'nombre': m.trabajador.get_full_name(),
                'rut':   p.rut if p else 'S/I',
                'cargo': p.cargo if p else 'S/I',
                'empresa': empresa.nombre,
                'entrada': None, 'ini_col': None,
                'fin_col': None, 'salida':  None,
            }
        h = fl.time()
        if m.tipo == 'ENTRADA':
            if reporte[key]['entrada'] is None or h < reporte[key]['entrada']:
                reporte[key]['entrada'] = h
        elif m.tipo == 'INICIO_COLACION': reporte[key]['ini_col'] = h
        elif m.tipo == 'FIN_COLACION':   reporte[key]['fin_col']  = h
        elif m.tipo == 'SALIDA':
            if reporte[key]['salida'] is None or h > reporte[key]['salida']:
                reporte[key]['salida'] = h

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Asistencia'

    headers = ['Fecha','Empresa','Trabajador','RUT','Cargo',
               'Entrada','Ini.Col','Fin.Col','T.Col','Salida','Horas']
    ws.append(headers)
    hf = PatternFill(start_color='1a1a1a', end_color='1a1a1a', fill_type='solid')
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = hf

    dummy = datetime.datetime.min
    for d in reporte.values():
        e, s   = d['entrada'], d['salida']
        ic, fc = d['ini_col'], d['fin_col']
        col_s  = horas_s = ''
        if ic and fc:
            diff  = datetime.datetime.combine(dummy, fc) - datetime.datetime.combine(dummy, ic)
            col_s = f"{int(diff.total_seconds()//3600):02d}:{int((diff.total_seconds()%3600)//60):02d}"
        if e and s:
            diff   = datetime.datetime.combine(dummy, s) - datetime.datetime.combine(dummy, e)
            horas_s= f"{int(diff.total_seconds()//3600):02d}:{int((diff.total_seconds()%3600)//60):02d}"
        ws.append([
            d['fecha'].strftime('%d/%m/%Y'), d['empresa'],
            d['nombre'], d['rut'], d['cargo'],
            e.strftime('%H:%M')  if e  else '--',
            ic.strftime('%H:%M') if ic else '--',
            fc.strftime('%H:%M') if fc else '--',
            col_s,
            s.strftime('%H:%M')  if s  else '--',
            horas_s,
        ])

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 14
    ws.column_dimensions['C'].width = 26

    resp = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="Asistencia_{empresa.nombre}.xlsx"'
    wb.save(resp)
    return resp


@login_required
@requiere_empleador
@requiere_empresa
def exportar_clima(request):
    """Excel de clima laboral (ánimo en salidas)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Clima Laboral'

    headers = ['Fecha','Hora','Trabajador','RUT','Cargo','Ánimo','Comentario']
    ws.append(headers)
    hf = PatternFill(start_color='4F4F4F', end_color='4F4F4F', fill_type='solid')
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = hf

    marcas = (
        Marcacion.objects
        .filter(empresa=request.empresa, tipo='SALIDA', animo__isnull=False)
        .select_related('trabajador', 'trabajador__perfil')
        .order_by('-timestamp')
    )
    emojis = {'FELIZ': '😄', 'NEUTRAL': '😐', 'MOLESTO': '😔'}
    for m in marcas:
        fl = timezone.localtime(m.timestamp)
        p  = getattr(m.trabajador, 'perfil', None)
        ws.append([
            fl.strftime('%d/%m/%Y'), fl.strftime('%H:%M'),
            m.trabajador.get_full_name(),
            p.rut if p else 'S/I',
            p.cargo if p else 'S/I',
            f"{emojis.get(m.animo,'')} {m.get_animo_display()}",
            m.comentario_animo or '-',
        ])

    ws.column_dimensions['C'].width = 26
    ws.column_dimensions['G'].width = 50

    resp = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = 'attachment; filename="Clima_Laboral.xlsx"'
    wb.save(resp)
    return resp


@login_required
def exportar_reporte_fiscalizacion(request):
    """Excel oficial Art. 33 CT con hash SHA-256 por registro."""
    perfil = getattr(request.user, 'perfil', None)
    if not perfil or not perfil.empresa:
        return HttpResponse('Sin empresa.', status=403)

    empresa = perfil.empresa
    desde   = request.GET.get('desde', '')
    hasta   = request.GET.get('hasta', '')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Libro Asistencia DT'

    hf   = PatternFill(start_color='003366', end_color='003366', fill_type='solid')
    hwf  = Font(bold=True, color='FFFFFF')
    ca   = Alignment(horizontal='center')

    headers = ['ID','RUT','Nombre','Fecha','Hora','Tipo','Origen','Dirección','Estado','Hash SHA-256']
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = hwf; cell.fill = hf; cell.alignment = ca

    marcas = (
        Marcacion.objects
        .filter(empresa=empresa)
        .select_related('trabajador', 'trabajador__perfil')
        .order_by('-timestamp')
    )
    if desde and hasta:
        marcas = marcas.filter(timestamp__range=[desde, hasta + ' 23:59:59'])

    for r, m in enumerate(marcas, 2):
        fl  = timezone.localtime(m.timestamp)
        p   = getattr(m.trabajador, 'perfil', None)
        rut = p.rut if p else 'S/I'
        sha = hashlib.sha256(
            f"{m.id}{rut}{fl.strftime('%d/%m/%Y')}{fl.strftime('%H:%M:%S')}{m.tipo}".encode()
        ).hexdigest()

        ws.cell(r, 1, m.id).alignment = ca
        ws.cell(r, 2, rut).alignment  = ca
        ws.cell(r, 3, m.trabajador.get_full_name())
        ws.cell(r, 4, fl.strftime('%d/%m/%Y')).alignment = ca
        ws.cell(r, 5, fl.strftime('%H:%M:%S')).alignment = ca
        ct = ws.cell(r, 6, m.tipo)
        ct.alignment = ca
        if m.tipo == 'ENTRADA': ct.font = Font(color='006600', bold=True)
        elif m.tipo == 'SALIDA': ct.font = Font(color='990000', bold=True)
        ws.cell(r, 7, 'WEB/APP').alignment = ca
        ws.cell(r, 8, (m.direccion or f'{m.latitud},{m.longitud}')[:60])
        ws.cell(r, 9, m.estado).alignment = ca
        ws.cell(r, 10, sha).font = Font(name='Courier New', size=9, color='555555')

    widths = [12,15,28,12,12,16,10,40,12,66]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    resp = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="LibroAsistencia_DT_{desde}_{hasta}.xlsx"'
    wb.save(resp)
    return resp


@login_required
@requiere_empleador
@requiere_empresa
def exportar_remuneraciones(request):
    """Pre-nómina LRE: días, H.E. 50%, H.E. 100%, atrasos, ausencias."""
    empresa = request.empresa
    hoy     = timezone.localdate()
    fi_str  = request.GET.get('fecha_inicio')
    ff_str  = request.GET.get('fecha_fin')

    if fi_str and ff_str:
        start = datetime.datetime.strptime(fi_str, '%Y-%m-%d').date()
        end   = datetime.datetime.strptime(ff_str, '%Y-%m-%d').date()
    else:
        start = hoy.replace(day=1)
        nm    = start.replace(day=28) + timedelta(days=4)
        end   = nm - timedelta(days=nm.day)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Pre-Nómina LRE'

    hf  = PatternFill(start_color='2F75B5', end_color='2F75B5', fill_type='solid')
    hwf = Font(color='FFFFFF', bold=True)
    ca  = Alignment(horizontal='center')

    headers = ['RUT','Nombre','Cargo','Días Trab.',
               'H.Ord (HH:MM)','H.E.50% (HH:MM)','H.E.100% (HH:MM)',
               'Min.Atraso','Días Ausencia','Observaciones']
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = hf; cell.font = hwf; cell.alignment = ca

    feriados = set(
        Feriado.objects.filter(fecha__range=[start, end]).values_list('fecha', flat=True)
    )

    from django.contrib.auth.models import User as DjangoUser
    trabajadores = DjangoUser.objects.filter(
        perfil__empresa=empresa, is_active=True
    ).select_related('perfil')

    def fmt(s):
        return f"{int(s//3600):02d}:{int((s%3600)//60):02d}"

    for trabajador in trabajadores:
        p = getattr(trabajador, 'perfil', None)
        if not p: continue

        jornada_s  = (p.jornada_diaria or 9) * 3600
        h_entrada  = p.hora_entrada or time(9, 0)
        tolerancia = empresa.tolerancia_atraso_min

        dias_trab = seg_ord = seg_e50 = seg_e100 = min_atraso = dias_aus = 0
        obs = []

        for i in range((end - start).days + 1):
            dia     = start + timedelta(days=i)
            dia_sem = dia.weekday()
            mapa    = {0: p.trabaja_lunes, 1: p.trabaja_martes, 2: p.trabaja_miercoles,
                       3: p.trabaja_jueves, 4: p.trabaja_viernes,
                       5: p.trabaja_sabado, 6: p.trabaja_domingo}
            le_toca = mapa.get(dia_sem, False)

            es_feriado = dia in feriados
            es_dom     = dia_sem == 6
            es_sab     = dia_sem == 5
            es_vac     = Vacacion.objects.filter(
                trabajador=trabajador, inicio__lte=dia, fin__gte=dia, estado='APROBADA'
            ).exists()
            es_lic     = LicenciaMedica.objects.filter(
                trabajador=trabajador, inicio__lte=dia, fin__gte=dia
            ).exists()

            marcas_dia = Marcacion.objects.filter(trabajador=trabajador, timestamp__date=dia)

            if marcas_dia.exists():
                dias_trab += 1
                ent = marcas_dia.filter(tipo='ENTRADA').aggregate(Min('timestamp'))['timestamp__min']
                sal = marcas_dia.filter(tipo='SALIDA').aggregate(Max('timestamp'))['timestamp__max']
                if ent and sal:
                    neto = max(0, (sal - ent).total_seconds() - 3600)
                    if es_dom or es_feriado:
                        seg_e100 += neto
                    elif neto > jornada_s:
                        seg_ord  += jornada_s
                        seg_e50  += neto - jornada_s
                    else:
                        seg_ord  += neto

                    if not (es_feriado or es_dom or es_sab) and ent:
                        hr  = timezone.localtime(ent).time()
                        dum = datetime.datetime.today().date()
                        dt_r = datetime.datetime.combine(dum, hr)
                        dt_o = datetime.datetime.combine(dum, h_entrada)
                        if dt_r > dt_o + timedelta(minutes=tolerancia):
                            min_atraso += int((dt_r - dt_o).total_seconds() / 60)
            else:
                if not (es_feriado or es_vac or es_lic or not le_toca):
                    dias_aus += 1

        if dias_aus > 0:
            obs.append(f'{dias_aus} ausencia(s)')
        if es_vac:
            obs.append('Con vacaciones')

        ws.append([
            p.rut or 'S/I', trabajador.get_full_name(), p.cargo or 'S/I',
            dias_trab, fmt(seg_ord), fmt(seg_e50), fmt(seg_e100),
            min_atraso, dias_aus, ', '.join(obs),
        ])

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18

    resp = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="PreNomina_{empresa.nombre}_{start}.xlsx"'
    wb.save(resp)
    return resp


@login_required
def generar_pdf_trabajador(request):
    """PDF de libro de asistencia personal usando ReportLab."""
    from io import BytesIO
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT

    user    = request.user
    marcas  = Marcacion.objects.filter(trabajador=user).order_by('timestamp')
    empresa = getattr(getattr(user, 'perfil', None), 'empresa', None)
    perfil  = getattr(user, 'perfil', None)

    buffer = BytesIO()
    doc    = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=1.5*cm, leftMargin=1.5*cm,
        topMargin=1.5*cm,   bottomMargin=1.5*cm,
    )

    styles  = getSampleStyleSheet()
    story   = []

    # Título
    titulo_style = ParagraphStyle(
        'titulo', parent=styles['Title'],
        fontSize=14, spaceAfter=6, alignment=TA_CENTER,
    )
    sub_style = ParagraphStyle(
        'sub', parent=styles['Normal'],
        fontSize=9, spaceAfter=4, alignment=TA_CENTER, textColor=colors.grey,
    )

    story.append(Paragraph('LIBRO DE ASISTENCIA DIGITAL', titulo_style))
    story.append(Paragraph(
        f'{empresa.nombre if empresa else ""} · Art. 33 Código del Trabajo Chile',
        sub_style
    ))
    story.append(Spacer(1, 0.3*cm))

    # Info del trabajador
    info_data = [
        ['Trabajador:', user.get_full_name(), 'RUT:', perfil.rut if perfil else '–'],
        ['Empresa:', empresa.nombre if empresa else '–', 'Cargo:', perfil.cargo if perfil else '–'],
        ['Generado:', timezone.localtime(timezone.now()).strftime('%d/%m/%Y %H:%M'), 'Total registros:', str(marcas.count())],
    ]
    info_table = Table(info_data, colWidths=[3*cm, 7*cm, 3*cm, 7*cm])
    info_table.setStyle(TableStyle([
        ('FONTSIZE',    (0,0), (-1,-1), 8),
        ('FONTNAME',    (0,0), (0,-1), 'Helvetica-Bold'),
        ('FONTNAME',    (2,0), (2,-1), 'Helvetica-Bold'),
        ('TEXTCOLOR',   (0,0), (0,-1), colors.grey),
        ('TEXTCOLOR',   (2,0), (2,-1), colors.grey),
        ('BOTTOMPADDING',(0,0),(-1,-1), 3),
        ('TOPPADDING',  (0,0),(-1,-1), 3),
    ]))
    story.append(info_table)
    story.append(Spacer(1, 0.4*cm))

    # Tabla de marcaciones
    headers = ['Fecha', 'Hora', 'Tipo', 'Ubicación', 'Manual', 'Hash SHA-256']
    data    = [headers]

    tipo_labels = {
        'ENTRADA': 'Entrada', 'SALIDA': 'Salida',
        'INICIO_COLACION': 'Ini. Colación', 'FIN_COLACION': 'Fin Colación',
    }

    for m in marcas:
        fl = timezone.localtime(m.timestamp)
        data.append([
            fl.strftime('%d/%m/%Y'),
            fl.strftime('%H:%M:%S'),
            tipo_labels.get(m.tipo, m.tipo),
            (m.direccion or '–')[:50],
            'Sí' if m.es_manual else 'No',
            m.hash_actual[:32] + '...' if m.hash_actual else '–',
        ])

    col_widths = [2.5*cm, 2.2*cm, 3*cm, 8*cm, 1.5*cm, 8*cm]
    tabla = Table(data, colWidths=col_widths, repeatRows=1)
    tabla.setStyle(TableStyle([
        # Header
        ('BACKGROUND',   (0,0), (-1,0), colors.HexColor('#1a1a1a')),
        ('TEXTCOLOR',    (0,0), (-1,0), colors.white),
        ('FONTNAME',     (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE',     (0,0), (-1,0), 8),
        ('ALIGN',        (0,0), (-1,0), 'CENTER'),
        # Datos
        ('FONTSIZE',     (0,1), (-1,-1), 7),
        ('FONTNAME',     (0,1), (-1,-1), 'Helvetica'),
        ('ROWBACKGROUNDS',(0,1),(-1,-1), [colors.white, colors.HexColor('#f7f7f5')]),
        ('GRID',         (0,0), (-1,-1), 0.25, colors.HexColor('#e0e0e0')),
        ('TOPPADDING',   (0,0), (-1,-1), 3),
        ('BOTTOMPADDING',(0,0), (-1,-1), 3),
        ('LEFTPADDING',  (0,0), (-1,-1), 4),
    ]))
    story.append(tabla)

    doc.build(story)
    buffer.seek(0)

    resp = HttpResponse(buffer, content_type='application/pdf')
    resp['Content-Disposition'] = f'attachment; filename="asistencia_{user.username}.pdf"'
    return resp
